from io import BytesIO
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pydantic import BaseModel
import ollama
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pathlib import Path
import fitz
import uuid
import json

# --------------------------------------------------
# 1. Create FastAPI Application
# --------------------------------------------------

app = FastAPI(
    title="Government Document Analysis System",
    description="Offline AI-Powered Government Document Analysis and Excel Automation",
    version="1.0.0"
)


# --------------------------------------------------
# 2. Configure CORS
# --------------------------------------------------
class DocumentAnalysisRequest(BaseModel):
    document_text: str

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:4200",
        "http://127.0.0.1:4200"
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"]
)


# --------------------------------------------------
# 3. Configure Upload Directory
# --------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent

UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

MAX_FILE_SIZE = 20 * 1024 * 1024  # 20 MB per PDF


# --------------------------------------------------
# 4. Home Endpoint
# --------------------------------------------------

@app.get("/")
def home():
    return {
        "message": "Government Document Analysis System API is running!",
        "status": "success"
    }


# --------------------------------------------------
# 5. Health Check Endpoint
# --------------------------------------------------

@app.get("/health")
def health_check():
    return {
        "status": "healthy",
        "application": "GovDocAI"
    }


# --------------------------------------------------
# 6. Upload and Validate PDF Documents
# --------------------------------------------------

@app.post("/upload")
async def upload_documents(
    files: list[UploadFile] = File(...)
):

    if not files:
        raise HTTPException(
            status_code=400,
            detail="Please upload at least one PDF document."
        )

    results = []

    for file in files:

        filename = file.filename or "unknown.pdf"

        try:

            # Validate the file extension.
            if Path(filename).suffix.lower() != ".pdf":
                results.append({
                    "filename": filename,
                    "status": "error",
                    "message": "Only PDF files are allowed."
                })
                continue

            # Read the uploaded file.
            content = await file.read()

            # Reject empty files.
            if not content:
                results.append({
                    "filename": filename,
                    "status": "error",
                    "message": "The uploaded file is empty."
                })
                continue

            # Check the file size.
            if len(content) > MAX_FILE_SIZE:
                results.append({
                    "filename": filename,
                    "status": "error",
                    "message": "File exceeds the 20 MB limit."
                })
                continue

            # Extract text and validate the PDF.
            with fitz.open(
                stream=content,
                filetype="pdf"
            ) as pdf:

                page_count = len(pdf)

                if page_count == 0:
                    results.append({
                        "filename": filename,
                        "status": "error",
                        "message": "The PDF contains no pages."
                    })
                    continue

                extracted_text = "\n".join(
                    page.get_text()
                    for page in pdf
                ).strip()

            # Generate a unique filename.
            saved_name = f"{uuid.uuid4().hex}.pdf"

            saved_path = UPLOAD_DIR / saved_name

            # Save the original PDF.
            saved_path.write_bytes(content)

            # Add the result.
            results.append({
                "filename": filename,
                "status": "success",
                "saved_as": saved_name,
                "pages": page_count,
                "text_extracted": bool(extracted_text),
                "text_length": len(extracted_text),
                "message": "PDF uploaded and validated successfully."
            })

        except Exception:
            results.append({
                "filename": filename,
                "status": "error",
                "message": "The file could not be processed as a valid PDF."
            })

        finally:
            await file.close()

    # Calculate processing statistics.
    successful_count = sum(
        1
        for result in results
        if result["status"] == "success"
    )

    failed_count = len(results) - successful_count

    return {
        "message": "Document upload completed.",
        "total_received": len(files),
        "successful_uploads": successful_count,
        "failed_uploads": failed_count,
        "files": results
    }

@app.post("/analyze")
def analyze_document(request: DocumentAnalysisRequest):

    if not request.document_text.strip():
        raise HTTPException(
            status_code=400,
            detail="Document text cannot be empty."
        )

    try:
        response = ollama.chat(
            model="llama3.2:3b",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are a government document analysis assistant. "
                        "Extract information only from the provided text. "
                        "Do not invent missing information. "
                        "Return valid JSON containing company_name, "
                        "document_type, annual_turnover, net_profit, "
                        "udyam_registration_number, tender_reference_number, "
                        "and tender_value. Use null for unavailable fields."
                    )
                },
                {
                    "role": "user",
                    "content": request.document_text
                }
            ],
            format="json"
        )

        return {
            "status": "success",
            "analysis": response["message"]["content"]
        }

    except Exception as error:
        raise HTTPException(
            status_code=500,
            detail=f"AI analysis failed: {str(error)}"
        )
# --------------------------------------------------
# 7. Automatic AI-Powered Document Processing
# --------------------------------------------------

@app.post("/process-documents")
async def process_documents(
    files: list[UploadFile] = File(...)
):
    if not files:
        raise HTTPException(
            status_code=400,
            detail="Please upload PDF documents."
        )

    results = []

    for file in files:
        filename = file.filename or "unknown.pdf"

        # Validate the file extension.
        if Path(filename).suffix.lower() != ".pdf":
            results.append({
                "filename": filename,
                "status": "error",
                "message": "Only PDF files are allowed."
            })
            continue

        try:
            # Read the PDF.
            content = await file.read()

            if not content:
                results.append({
                    "filename": filename,
                    "status": "error",
                    "message": "The file is empty."
                })
                continue

            if len(content) > MAX_FILE_SIZE:
                results.append({
                    "filename": filename,
                    "status": "error",
                    "message": "File exceeds the 20 MB limit."
                })
                continue

            # Extract text from the PDF.
            with fitz.open(
                stream=content,
                filetype="pdf"
            ) as pdf:
                page_count = len(pdf)

                if page_count == 0:
                    results.append({
                        "filename": filename,
                        "status": "error",
                        "message": "The PDF contains no pages."
                    })
                    continue

                document_text = "\n".join(
                    page.get_text()
                    for page in pdf
                ).strip()

            if not document_text:
                results.append({
                    "filename": filename,
                    "status": "error",
                    "message": (
                        "No readable text was found. "
                        "OCR may be required for scanned PDFs."
                    )
                })
                continue

            # Ask the local Ollama model to analyze the document.
            response = ollama.chat(
                model="llama3.2:3b",
                messages=[
                    {
                        "role": "system",
                        "content": """
You analyze government company documents.

Determine whether the document is:
- Financial Document
- MSME Document
- Tender Document
- Unknown

Extract information only when supported by the document.

Return valid JSON with these fields:
{
  "company_name": null,
  "document_type": "Unknown",
  "financial_year": null,
  "annual_turnover": null,
  "net_profit": null,
  "net_worth": null,
  "udyam_registration_number": null,
  "enterprise_type": null,
  "tender_reference_number": null,
  "tender_value": null,
  "bid_deadline": null
}

Use null for information that is not available.
Do not invent values.
"""
                    },
                    {
                        "role": "user",
                        "content": (
                            "Analyze this document:\n\n"
                            + document_text[:12000]
                        )
                    }
                ],
                format="json"
            )

            # Convert the model response into a JSON object.
            import json

            analysis_text = response["message"]["content"]
            analysis_data = json.loads(analysis_text)

            results.append({
                "filename": filename,
                "status": "success",
                "pages": page_count,
                "analysis": analysis_data
            })

        except json.JSONDecodeError:
            results.append({
                "filename": filename,
                "status": "error",
                "message": "The AI returned invalid JSON."
            })

        except Exception as error:
            results.append({
                "filename": filename,
                "status": "error",
                "message": (
                    "Document analysis failed: "
                    + str(error)
                )
            })

        finally:
            await file.close()

    successful_count = sum(
        1
        for result in results
        if result["status"] == "success"
    )

    return {
        "message": "Document processing completed.",
        "total_received": len(files),
        "successful_analyses": successful_count,
        "failed_analyses": len(files) - successful_count,
        "documents": results
    }

# --------------------------------------------------
# 8. Consolidate Documents by Company
# --------------------------------------------------

def consolidate_company_documents(documents):
    companies = {}

    for document in documents:
        if document.get("status") != "success":
            continue

        analysis = document.get("analysis", {})

        company_name = analysis.get("company_name")

        if not company_name:
            continue

        company_name = company_name.strip()

        # Create a company record if it does not exist.
        if company_name not in companies:
            companies[company_name] = {
                "company_name": company_name,
                "financial_year": None,
                "annual_turnover": None,
                "net_profit": None,
                "net_worth": None,
                "udyam_registration_number": None,
                "enterprise_type": None,
                "tender_reference_number": None,
                "tender_value": None,
                "bid_deadline": None,
                "documents_received": [],
                "review_required": False
            }

        company = companies[company_name]

        document_type = (
            analysis.get("document_type") or ""
        ).lower()

        # Track the source documents.
        if document.get("filename"):
            company["documents_received"].append(
                document["filename"]
            )

        # Populate financial fields only from financial documents.
        if (
            "financial" in document_type
            or "annual financial statement" in document_type
        ):
            financial_fields = [
                "financial_year",
                "annual_turnover",
                "net_profit",
                "net_worth"
            ]

            for field in financial_fields:
                value = analysis.get(field)

                if value is not None and value != "null":
                    company[field] = value

        # Populate MSME fields only from MSME documents.
        elif "msme" in document_type:
            msme_fields = [
                "udyam_registration_number",
                "enterprise_type"
            ]

            for field in msme_fields:
                value = analysis.get(field)

                if value is not None and value != "null":
                    company[field] = value

        # Populate tender fields only from tender documents.
        elif "tender" in document_type:
            tender_fields = [
                "tender_reference_number",
                "tender_value",
                "bid_deadline"
            ]

            for field in tender_fields:
                value = analysis.get(field)

                if value is not None and value != "null":
                    company[field] = value

        else:
            company["review_required"] = True

    # Determine whether any required information is missing.
    for company in companies.values():
        required_fields = [
            "annual_turnover",
            "udyam_registration_number",
            "tender_reference_number"
        ]

        if any(
            company[field] is None
            for field in required_fields
        ):
            company["review_required"] = True

    return list(companies.values())
# --------------------------------------------------
# 9. Company Consolidation API
# --------------------------------------------------

from pydantic import BaseModel


class ConsolidationRequest(BaseModel):
    documents: list[dict]


@app.post("/consolidate")
def consolidate_documents(request: ConsolidationRequest):

    company_records = consolidate_company_documents(
        request.documents
    )

    return {
        "message": "Company document consolidation completed.",
        "companies_found": len(company_records),
        "companies": company_records
    }

# --------------------------------------------------
# 10. Generate Excel Report
# --------------------------------------------------

@app.post("/export-excel")
def export_excel(request: ConsolidationRequest):

    # Consolidate the supplied document analysis results.
    company_records = consolidate_company_documents(
        request.documents
    )

    if not company_records:
        raise HTTPException(
            status_code=400,
            detail="No company records available for export."
        )

    # Create the Excel workbook.
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Company Analysis"

    # Define the Excel columns.
    columns = [
        "Company Name",
        "Financial Year",
        "Annual Turnover",
        "Net Profit",
        "Net Worth",
        "Udyam Registration Number",
        "Enterprise Type",
        "Tender Reference Number",
        "Tender Value",
        "Bid Deadline",
        "Documents Received",
        "Review Required"
    ]

    worksheet.append(columns)

    # Style the header row.
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Add one row for each company.
    for company in company_records:
        worksheet.append([
            company.get("company_name"),
            company.get("financial_year"),
            company.get("annual_turnover"),
            company.get("net_profit"),
            company.get("net_worth"),
            company.get("udyam_registration_number"),
            company.get("enterprise_type"),
            company.get("tender_reference_number"),
            company.get("tender_value"),
            company.get("bid_deadline"),
            ", ".join(company.get("documents_received", [])),
            "Yes" if company.get("review_required") else "No"
        ])

    # Adjust column widths.
    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )

        column_letter = column_cells[0].column_letter

        worksheet.column_dimensions[column_letter].width = min(
            max_length + 3,
            40
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    # Save the workbook in memory.
    excel_buffer = BytesIO()
    workbook.save(excel_buffer)
    excel_buffer.seek(0)

    return StreamingResponse(
        excel_buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                'attachment; filename="GovDocAI_Report.xlsx"'
        }
    )

# --------------------------------------------------
# 11. Complete Automated Processing and Excel Export
# --------------------------------------------------

@app.post("/process-and-export")
async def process_and_export(
    files: list[UploadFile] = File(...)
):
    if not files:
        raise HTTPException(
            status_code=400,
            detail="Please upload PDF documents."
        )

    analyzed_documents = []

    for file in files:
        filename = file.filename or "unknown.pdf"

        try:
            # Validate file type.
            if Path(filename).suffix.lower() != ".pdf":
                raise ValueError("Only PDF files are allowed.")

            # Read PDF contents.
            content = await file.read()

            if not content:
                raise ValueError("The uploaded PDF is empty.")

            if len(content) > MAX_FILE_SIZE:
                raise ValueError("PDF exceeds the 20 MB limit.")

            # Extract PDF text.
            with fitz.open(
                stream=content,
                filetype="pdf"
            ) as pdf:

                if len(pdf) == 0:
                    raise ValueError("The PDF contains no pages.")

                page_count = len(pdf)

                document_text = "\n".join(
                    page.get_text()
                    for page in pdf
                ).strip()

            if not document_text:
                raise ValueError(
                    "No readable text found. OCR may be required."
                )

            # Analyze the document using local Ollama.
            response = ollama.chat(
                model="llama3.2:3b",
                messages=[
                    {
                        "role": "system",
                        "content": """
You are a government document analysis assistant.

Classify the document as exactly one of:
Financial Document, MSME Document, Tender Document, Unknown.

Extract information only from the provided document.
Never invent values.
Use JSON null when information is unavailable.

Return valid JSON with these fields:
{
  "company_name": null,
  "document_type": "Unknown",
  "financial_year": null,
  "annual_turnover": null,
  "net_profit": null,
  "net_worth": null,
  "udyam_registration_number": null,
  "enterprise_type": null,
  "tender_reference_number": null,
  "tender_value": null,
  "bid_deadline": null
}

Use financial values only when explicitly supported
by the document. Do not copy financial values into
unrelated document types.
"""
                    },
                    {
                        "role": "user",
                        "content": (
                            "Analyze this document:\n\n"
                            + document_text[:12000]
                        )
                    }
                ],
                format="json"
            )

            analysis = json.loads(
                response["message"]["content"]
            )

            analyzed_documents.append({
                "filename": filename,
                "status": "success",
                "pages": page_count,
                "analysis": analysis
            })

        except Exception as error:
            analyzed_documents.append({
                "filename": filename,
                "status": "error",
                "message": str(error)
            })

        finally:
            await file.close()

    # Consolidate successful AI results.
    company_records = consolidate_company_documents(
        analyzed_documents
    )

    if not company_records:
        raise HTTPException(
            status_code=422,
            detail={
                "message": "No company records could be created.",
                "documents": analyzed_documents
            }
        )

    # Create the Excel workbook.
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Company Analysis"

    columns = [
        "Company Name",
        "Financial Year",
        "Annual Turnover",
        "Net Profit",
        "Net Worth",
        "Udyam Registration Number",
        "Enterprise Type",
        "Tender Reference Number",
        "Tender Value",
        "Bid Deadline",
        "Documents Received",
        "Review Required"
    ]

    worksheet.append(columns)

    # Style the header.
    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Write one row per company.
    for company in company_records:
        worksheet.append([
            company.get("company_name"),
            company.get("financial_year"),
            company.get("annual_turnover"),
            company.get("net_profit"),
            company.get("net_worth"),
            company.get("udyam_registration_number"),
            company.get("enterprise_type"),
            company.get("tender_reference_number"),
            company.get("tender_value"),
            company.get("bid_deadline"),
            ", ".join(
                company.get("documents_received", [])
            ),
            "Yes" if company.get("review_required") else "No"
        ])

    # Adjust column widths.
    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )

        column_letter = column_cells[0].column_letter

        worksheet.column_dimensions[
            column_letter
        ].width = min(max_length + 3, 40)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    # Prepare Excel file for download.
    excel_buffer = BytesIO()
    workbook.save(excel_buffer)
    excel_buffer.seek(0)

    return StreamingResponse(
        excel_buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                'attachment; filename="GovDocAI_Report.xlsx"'
        }
    )